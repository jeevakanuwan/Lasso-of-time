from pathlib import Path
import re, math, random, hashlib
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path("."); DRAFTS=ROOT/"06_drafts"; OUTDIR=ROOT/"proofreading"; ART=OUTDIR/"art"
OUTDIR.mkdir(exist_ok=True); ART.mkdir(exist_ok=True)
chapters=sorted(DRAFTS.glob("chapter_*.md"),key=lambda p:int(re.search(r"chapter_(\d+)",p.stem).group(1)))
if len(chapters)!=24: raise SystemExit(f"Expected 24 chapter drafts, found {len(chapters)}")

def font(size,bold=False,serif=False):
    opts=(["/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"] if serif else [])+["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    for f in opts:
        if Path(f).exists(): return ImageFont.truetype(f,size)
    return ImageFont.load_default()

def make_art(path,title,index=0,cover=False):
    W,H=(1500,1900) if cover else (1500,300)
    rng=random.Random(int(hashlib.sha256((title+str(index)).encode()).hexdigest()[:8],16))
    im=Image.new("RGB",(W,H),(5,10,28)); px=im.load()
    for y in range(H):
        for x in range(W):
            glow=max(0,1-math.sqrt(((x-W*.52)/(W*.82))**2+((y-H*.45)/(H*.8))**2))
            px[x,y]=(int(5+7*glow),int(10+19*glow),int(28+37*glow))
    d=ImageDraw.Draw(im,"RGBA")
    for _ in range(240 if cover else 95):
        x,y=rng.randrange(W),rng.randrange(H); r=rng.choice([1,1,2,3])
        d.ellipse((x-r,y-r,x+r,y+r),fill=(185,213,255,rng.randrange(90,220)))
    cx,cy=(W*.52,H*.40) if cover else (W*.78,H*.46); rad=H*.18 if cover else H*.48
    d.ellipse((cx-rad,cy-rad,cx+rad,cy+rad),fill=(20,70,135,180),outline=(78,169,255,230),width=5 if cover else 3)
    d.ellipse((cx-rad*.75,cy-rad*.9,cx+rad*.75,cy+rad*.9),outline=(104,208,255,190),width=3)
    d.arc((cx-rad*.95,cy-rad*.42,cx+rad*.95,cy+rad*.42),0,360,fill=(105,193,230,160),width=3)
    pts=[]
    for i in range(420):
        t=i/419*math.pi*2
        pts.append((W*.50+math.cos(t)*W*.29,H*.42+math.sin(t)*H*(.105 if cover else .24)+math.sin(t*2)*H*.018))
    d.line(pts,fill=(245,190,83,235),width=7 if cover else 4)
    d.ellipse((W*.50-8,H*.42-8,W*.50+8,H*.42+8),fill=(255,229,160,255))
    if cover:
        d.rectangle((0,int(H*.72),W,H),fill=(3,7,18,255))
        for _ in range(48):
            x=rng.randrange(W); bw=rng.randrange(12,48); bh=rng.randrange(30,140)
            d.rectangle((x,H*.72-bh,x+bw,H*.72),fill=(8,19,37,255))
        d.text((W*.10,H*.09),"A SCIENCE-FICTION NOVEL",font=font(30),fill=(158,193,226,255))
        d.text((W*.10,H*.83),"TIME  •  MEMORY  •  SURVIVAL",font=font(28,True),fill=(245,190,83,255))
    else:
        d.line([(65,H*.78),(W-65,H*.78)],fill=(79,120,174,100),width=1)
        for i in range(55):
            x=80+i*(W-160)/54; yy=H*.78+math.sin(i*.39+index)*H*.12
            d.ellipse((x-2,yy-2,x+2,yy+2),fill=(105,208,255,145))
    im.save(path)

def add_field(run,instr):
    a=OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"),"begin")
    b=OxmlElement("w:instrText"); b.set(qn("xml:space"),"preserve"); b.text=instr
    c=OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"),"end")
    run._r.append(a); run._r.append(b); run._r.append(c)

def add_inline(p,text):
    text=re.sub(r"!\[([^\]]*)\]\([^)]+\)",r"\1",text)
    text=re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",text)
    pat=re.compile(r"(\*\*.+?\*\*|__.+?__|\*.+?\*|_.+?_|\x60[^\x60]+\x60)")
    pos=0
    for m in pat.finditer(text):
        if m.start()>pos: p.add_run(text[pos:m.start()])
        tok=m.group()
        if tok.startswith("**") or tok.startswith("__"):
            r=p.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith(chr(96)):
            r=p.add_run(tok[1:-1]); r.font.name="Consolas"; r.font.size=Pt(9)
        else:
            r=p.add_run(tok[1:-1]); r.italic=True
        pos=m.end()
    if pos<len(text): p.add_run(text[pos:])

def add_block(doc,block):
    if not block:return
    if block.startswith("### "): doc.add_heading(block[4:].strip(),level=3); return
    if block.startswith("## "): doc.add_heading(block[3:].strip(),level=2); return
    if block.startswith(">"):
        p=doc.add_paragraph(style="Quote"); p.paragraph_format.left_indent=Inches(.22)
        add_inline(p,"\n".join(re.sub(r"^\s*>\s?","",x) for x in block.splitlines())); return
    if all(re.match(r"^\s*[-*+]\s+",x) for x in block.splitlines()):
        for x in block.splitlines():
            p=doc.add_paragraph(style="List Bullet"); add_inline(p,re.sub(r"^\s*[-*+]\s+","",x))
        return
    if all(re.match(r"^\s*\d+[.)]\s+",x) for x in block.splitlines()):
        for x in block.splitlines():
            p=doc.add_paragraph(style="List Number"); add_inline(p,re.sub(r"^\s*\d+[.)]\s+","",x))
        return
    p=doc.add_paragraph(); add_inline(p,re.sub(r"\s*\n\s*"," ",block))

cover=ART/"cover.png"; make_art(cover,"The Lasso of Time",cover=True)
arts=[]
for i,path in enumerate(chapters,1):
    lines=path.read_text(encoding="utf-8").splitlines()
    title=next((ln[2:].strip() for ln in lines if ln.startswith("# ")),"Chapter "+str(i))
    ap=ART/f"chapter_{i:02d}.png"; make_art(ap,title,i); arts.append(ap)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
sec.header_distance=Inches(.28); sec.footer_distance=Inches(.28)
styles=doc.styles; styles["Normal"].font.name="Georgia"; styles["Normal"].font.size=Pt(10.5); styles["Normal"].font.color.rgb=RGBColor(35,40,49)
styles["Normal"].paragraph_format.space_after=Pt(5); styles["Normal"].paragraph_format.line_spacing=1.08
for name,size,color in [("Title",32,(14,30,55)),("Heading 1",21,(14,45,80)),("Heading 2",15,(20,65,105)),("Heading 3",12,(25,75,110))]:
    st=styles[name]; st.font.name="Georgia"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(*color)
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=h.add_run("THE LASSO OF TIME  •  PROOFREADING EDITION"); r.font.name="Arial"; r.font.size=Pt(8); r.font.color.rgb=RGBColor(105,115,130)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run("Working copy — comments and Track Changes welcome  •  "); r.font.name="Arial"; r.font.size=Pt(8); r.font.color.rgb=RGBColor(105,115,130)
add_field(f.add_run(),"PAGE")

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10)
p.add_run().add_picture(str(cover),width=Inches(5.7))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("THE LASSO OF TIME"); r.font.name="Georgia"; r.font.size=Pt(30); r.bold=True; r.font.color.rgb=RGBColor(13,35,66)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("PROOFREADING EDITION"); r.font.name="Arial"; r.font.size=Pt(12); r.bold=True; r.font.color.rgb=RGBColor(174,125,39)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("Source: iteration/audit09102026  |  24 chapters"); r.font.name="Arial"; r.font.size=Pt(9); r.font.color.rgb=RGBColor(100,110,125)
doc.add_page_break()
doc.add_heading("Reader & proofreader notes",level=1)
doc.add_paragraph("This separate review copy is assembled from the prose drafts in 06_drafts/ on the GitHub branch iteration/audit09102026. Source chapter files are not edited by this build. Use Word comments and Review → Track Changes for suggestions; retain the source branch as the reference version.")
doc.add_paragraph("The original atmospheric chapter banners are decorative and do not assert additional story canon.")
doc.add_page_break()

for idx,path in enumerate(chapters,1):
    lines=path.read_text(encoding="utf-8").splitlines()
    title=next((ln[2:].strip() for ln in lines if ln.startswith("# ")),"Chapter "+str(idx))
    if idx>1: doc.add_page_break()
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    p.add_run().add_picture(str(arts[idx-1]),width=Inches(6.8))
    hd=doc.add_heading(title,level=1); hd.paragraph_format.space_before=Pt(2); hd.paragraph_format.space_after=Pt(8)
    buf=[]
    for ln in lines:
        if ln.startswith("# "): continue
        if not ln.strip():
            if buf: add_block(doc,"\n".join(buf).strip()); buf=[]
            continue
        if re.match(r"^\s*(---+|\*\*\*+|___+)\s*$",ln):
            if buf: add_block(doc,"\n".join(buf).strip()); buf=[]
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
            pPr=p._p.get_or_add_pPr(); borders=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
            bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"4"); bottom.set(qn("w:color"),"C7D5E5"); borders.append(bottom); pPr.append(borders)
            continue
        buf.append(ln)
    if buf:add_block(doc,"\n".join(buf).strip())

doc.core_properties.title="The Lasso of Time — Proofreading Edition"
doc.core_properties.subject="Editable proofreading copy generated from the audited GitHub branch"
doc.core_properties.author="Proofreading edition"
doc.core_properties.keywords="science fiction, proofreading, Track Changes"
out=OUTDIR/"The_Lasso_of_Time_Proofreading.docx"; doc.save(out)
print(f"Built {out} from {len(chapters)} chapters; source markdown files left unchanged.")
